#!/usr/bin/env python3
"""Enzo Crescendo deck — 9th-grade friendly language, no results."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

BG = RGBColor(0x14, 0x16, 0x1A)
PANEL = RGBColor(0x1E, 0x22, 0x28)
BRASS = RGBColor(0xC4, 0x96, 0x4A)
CREAM = RGBColor(0xF2, 0xEB, 0xDF)
MUTED = RGBColor(0xA8, 0xA0, 0x93)
LINE = RGBColor(0x3A, 0x34, 0x2C)

W, H = Inches(13.333), Inches(7.5)
SLIDE_W = 12192000
SLIDE_H = 6858000


def set_run(run, *, size=18, bold=False, color=CREAM, font="Georgia"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn("a:latin"))
    if latin is None:
        latin = rPr.makeelement(qn("a:latin"), {})
        rPr.insert(0, latin)
    latin.set("typeface", font)


def add_bg(slide, color=BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_accent_bar(slide):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BRASS
    bar.line.fill.background()


def add_footer(slide, page: int, total: int = 11):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(10), Inches(0.3))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f"ENZO  ·  FRC Crescendo  ·  Team 1577 Steampunk  ·  {page}/{total}"
    set_run(run, size=11, color=MUTED, font="Calibri")


def add_title_block(slide, kicker: str, title: str, subtitle: str | None = None):
    k = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(0.4))
    p = k.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = kicker.upper()
    set_run(run, size=13, bold=True, color=BRASS, font="Calibri")

    t = slide.shapes.add_textbox(Inches(0.7), Inches(0.85), Inches(12), Inches(1.2))
    tf = t.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    set_run(run, size=34, bold=True, color=CREAM, font="Georgia")

    if subtitle:
        s = slide.shapes.add_textbox(Inches(0.7), Inches(2.05), Inches(11.5), Inches(0.55))
        tf = s.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = subtitle
        set_run(run, size=16, color=MUTED, font="Calibri")


def add_bullets(slide, items: list[str], top=2.7, size=20):
    box = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.5), Inches(3.8))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(14)
        run = p.add_run()
        run.text = item
        set_run(run, size=size, color=CREAM, font="Calibri")


def add_card(slide, left, top, width, height, heading: str, body: str):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = PANEL
    card.line.color.rgb = LINE
    card.line.width = Pt(1)
    card.adjustments[0] = 0.08

    h = slide.shapes.add_textbox(
        Inches(left + 0.25), Inches(top + 0.25), Inches(width - 0.5), Inches(0.45)
    )
    p = h.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = heading
    set_run(run, size=18, bold=True, color=BRASS, font="Georgia")

    b = slide.shapes.add_textbox(
        Inches(left + 0.25), Inches(top + 0.85), Inches(width - 0.5), Inches(height - 1.15)
    )
    tf = b.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = body
    set_run(run, size=15, color=CREAM, font="Calibri")


def set_notes(slide, text: str):
    slide.notes_slide.notes_text_frame.text = text


def build() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_accent_bar(s)
    brand = s.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(11.5), Inches(1.2))
    p = brand.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "ENZO"
    set_run(run, size=72, bold=True, color=BRASS, font="Georgia")
    sub = s.shapes.add_textbox(Inches(0.8), Inches(3.2), Inches(11.5), Inches(1.0))
    tf = sub.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "Why Enzo is the best robot in Crescendo"
    set_run(run, size=28, color=CREAM, font="Georgia")
    meta = s.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(11.5), Inches(0.6))
    p = meta.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "Team 1577 Steampunk  ·  About smart design  ·  5–7 minutes"
    set_run(run, size=16, color=MUTED, font="Calibri")
    add_footer(s, 1)
    set_notes(s, "Say the robot name clearly. This talk is about design, not scores.")

    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_accent_bar(s)
    add_title_block(s, "Who is Enzo?", "Meet the robot")
    add_card(s, 0.7, 2.7, 3.8, 3.4, "The team", "Built by FRC Team 1577, Steampunk, from Raanana, Israel.")
    add_card(
        s,
        4.75,
        2.7,
        3.8,
        3.4,
        "The game",
        "Crescendo is like a sports match for robots: pick up foam rings called Notes, score them, and climb at the end.",
    )
    add_card(
        s,
        8.8,
        2.7,
        3.8,
        3.4,
        "Our focus",
        "We’ll talk about how Enzo was designed — not about game scores.",
    )
    add_footer(s, 2)
    set_notes(s, "Keep this short. Introduce the team and the game in plain words.")

    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_accent_bar(s)
    add_title_block(s, "The challenge", "What robots had to do")
    add_bullets(
        s,
        [
            "Pick up Notes from the floor quickly",
            "Shoot them into a big goal (the Speaker)",
            "Also put Notes in a side goal (the Amp)",
            "Climb onto a chain at the end of the match",
            "Keep moving even when other robots get in the way",
        ],
    )
    add_footer(s, 3)
    set_notes(s, "Explain Crescendo like explaining a new sport.")

    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_accent_bar(s)
    add_title_block(s, "What “best” means", "Not flashiest — smartest")
    add_bullets(
        s,
        [
            "Has a clear plan for the whole game",
            "Picks up and scores Notes smoothly",
            "Can help the team in more than one way",
            "Gets better when the team finds problems",
            "Is easy for the drivers to control",
        ],
    )
    add_footer(s, 4)
    set_notes(s, "These five points are our checklist for the rest of the talk.")

    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_accent_bar(s)
    add_title_block(
        s,
        "Improving the robot",
        "They weren’t afraid to change it",
        "Midway through the season, the team rebuilt big parts of Enzo.",
    )
    add_bullets(
        s,
        [
            "They listened to what wasn’t working and fixed it",
            "Great engineering means improving — not sticking with a bad first idea",
            "Enzo got better because the team was willing to change",
        ],
        top=2.9,
    )
    add_footer(s, 5)
    set_notes(s, "Praise the rebuild as smart teamwork, without talking about results.")

    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_accent_bar(s)
    add_title_block(s, "Picking up game pieces", "If you can’t grab Notes, you can’t score")
    add_card(
        s,
        0.7,
        2.7,
        5.9,
        3.5,
        "Simple idea",
        "Enzo is built around getting Notes fast and keeping them under control.",
    )
    add_card(
        s,
        6.9,
        2.7,
        5.7,
        3.5,
        "Sports comparison",
        "Owning the Notes is like owning the basketball. Enzo was made to take that job seriously.",
    )
    add_footer(s, 6)
    set_notes(s, "Use the basketball comparison if it helps the class.")

    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_accent_bar(s)
    add_title_block(s, "Scoring", "Two ways to score, one robot")
    add_bullets(
        s,
        [
            "It can shoot into the Speaker for regular scoring",
            "It can also score in the Amp to help its alliance",
            "Drivers can pick the better option when the field is crowded",
            "Being flexible is better than only being good at one thing",
        ],
    )
    add_footer(s, 7)
    set_notes(s, "Stress flexibility in simple terms.")

    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_accent_bar(s)
    add_title_block(s, "The end of the match", "Climbing without breaking the plan")
    add_card(
        s,
        0.7,
        2.7,
        5.9,
        3.5,
        "Common mistake",
        "Some robots try so hard to climb that they get worse at scoring.",
    )
    add_card(
        s,
        6.9,
        2.7,
        5.7,
        3.5,
        "Enzo’s balance",
        "Enzo stays strong at scoring during the match — and still can climb when it matters at the end.",
    )
    add_footer(s, 8)
    set_notes(s, "Balance: score first, climb later — but both matter.")

    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_accent_bar(s)
    add_title_block(s, "Driving the robot", "Good code makes hard jobs feel easy")
    add_bullets(
        s,
        [
            "Helps the robot aim and score more consistently",
            "Gives drivers clear controls instead of button chaos",
            "Uses cameras and sensors so the robot “sees” the field",
            "When drivers trust the robot, they can play smarter",
        ],
    )
    add_footer(s, 9)
    set_notes(s, "Software = helpful tools for the human drivers.")

    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_accent_bar(s)
    add_title_block(s, "The big idea", "Why Enzo is the best")
    reasons = [
        ("1", "Built for the whole Crescendo game, not just one trick"),
        ("2", "The team rebuilt it when they needed to"),
        ("3", "Focuses on grabbing and controlling Notes"),
        ("4", "Can score in more than one way"),
        ("5", "Software helps drivers do the hard parts"),
    ]
    for i, (n, text) in enumerate(reasons):
        y = 2.45 + i * 0.72
        num = s.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(0.75), Inches(y), Inches(0.48), Inches(0.48)
        )
        num.fill.solid()
        num.fill.fore_color.rgb = BRASS
        num.line.fill.background()
        nbox = s.shapes.add_textbox(Inches(0.75), Inches(y + 0.05), Inches(0.48), Inches(0.4))
        p = nbox.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = n
        set_run(run, size=14, bold=True, color=BG, font="Calibri")
        tbox = s.shapes.add_textbox(Inches(1.5), Inches(y + 0.05), Inches(10.8), Inches(0.5))
        p = tbox.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = text
        set_run(run, size=18, color=CREAM, font="Calibri")
    add_footer(s, 10)
    set_notes(s, "Slow down and hit each of the five points.")

    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_accent_bar(s)
    close = s.shapes.add_textbox(Inches(0.8), Inches(2.3), Inches(11.5), Inches(1.3))
    tf = close.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "Enzo is best because of how it was built."
    set_run(run, size=32, bold=True, color=CREAM, font="Georgia")
    sub = s.shapes.add_textbox(Inches(0.8), Inches(3.8), Inches(11.5), Inches(1.0))
    tf = sub.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "Clear plan. Willing to improve. Easy to drive. Ready for the whole game."
    set_run(run, size=18, color=MUTED, font="Calibri")
    q = s.shapes.add_textbox(Inches(0.8), Inches(5.3), Inches(11.5), Inches(0.5))
    p = q.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "Questions?"
    set_run(run, size=22, bold=True, color=BRASS, font="Georgia")
    add_footer(s, 11)
    set_notes(s, "End here. Invite simple questions about design or the game.")

    out = Path(__file__).resolve().parent / "Enzo_Crescendo_Presentation.pptx"
    prs.save(out)
    return out


if __name__ == "__main__":
    print(f"Wrote {build()}")
