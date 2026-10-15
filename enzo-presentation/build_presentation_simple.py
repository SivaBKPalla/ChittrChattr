#!/usr/bin/env python3
"""Simpler PPTX: 9th-grade language, Google Slides friendly."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

BG = RGBColor(20, 22, 26)
BRASS = RGBColor(196, 150, 74)
CREAM = RGBColor(242, 235, 223)
MUTED = RGBColor(168, 160, 147)

SLIDE_W = 12192000
SLIDE_H = 6858000


def style(run, size=20, bold=False, color=CREAM):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def fill_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_text(slide, left, top, width, height, text, *, size=20, bold=False, color=CREAM):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = line
        style(run, size=size, bold=bold, color=color)
    return box


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def build() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    slides = [
        {
            "kicker": "TEAM 1577 STEAMPUNK  ·  FRC CRESCENDO",
            "title": "ENZO",
            "body": "Why Enzo is the best robot in Crescendo\nA short talk about smart design  ·  about 5–7 minutes",
            "notes": "Say the robot name clearly. This talk is about design, not scores.",
        },
        {
            "kicker": "WHO IS ENZO?",
            "title": "Meet the robot",
            "body": "Enzo was built by FRC Team 1577, Steampunk, from Raanana, Israel.\n\nCrescendo is like a sports match for robots: pick up foam rings called Notes, score them, and climb at the end.\n\nWe’ll talk about how Enzo was designed — not about game scores.",
            "notes": "Keep this short. Introduce the team and the game in plain words.",
        },
        {
            "kicker": "THE CHALLENGE",
            "title": "What robots had to do",
            "body": "• Pick up Notes from the floor quickly\n• Shoot them into a big goal (the Speaker)\n• Also put Notes in a side goal (the Amp)\n• Climb onto a chain at the end of the match\n• Keep moving even when other robots get in the way",
            "notes": "Explain Crescendo like explaining a new sport.",
        },
        {
            "kicker": "WHAT “BEST” MEANS",
            "title": "Not flashiest — smartest",
            "body": "• Has a clear plan for the whole game\n• Picks up and scores Notes smoothly\n• Can help the team in more than one way\n• Gets better when the team finds problems\n• Is easy for the drivers to control",
            "notes": "These five points are our checklist for the rest of the talk.",
        },
        {
            "kicker": "IMPROVING THE ROBOT",
            "title": "They weren’t afraid to change it",
            "body": "Midway through the season, the team rebuilt big parts of Enzo.\n\nThat means they listened to what wasn’t working and fixed it.\n\nGreat engineering looks like that — not giving up on the first idea.",
            "notes": "Praise the rebuild as smart teamwork, without talking about results.",
        },
        {
            "kicker": "PICKING UP GAME PIECES",
            "title": "If you can’t grab Notes, you can’t score",
            "body": "Enzo is built around getting Notes fast and keeping them under control.\n\nIn this game, owning the Notes is like owning the basketball.\n\nEnzo was made to take that job seriously.",
            "notes": "Use the basketball comparison if it helps the class.",
        },
        {
            "kicker": "SCORING",
            "title": "Two ways to score, one robot",
            "body": "• It can shoot into the Speaker for regular scoring\n• It can also score in the Amp to help its alliance\n• Drivers can pick the better option when the field is crowded\n• Being flexible is better than only being good at one thing",
            "notes": "Stress flexibility in simple terms.",
        },
        {
            "kicker": "THE END OF THE MATCH",
            "title": "Climbing without breaking the plan",
            "body": "Some robots try so hard to climb that they get worse at scoring.\n\nEnzo stays strong at scoring during the match — and still can climb when it matters at the end.",
            "notes": "Balance: score first, climb later — but both matter.",
        },
        {
            "kicker": "DRIVING THE ROBOT",
            "title": "Good code makes hard jobs feel easy",
            "body": "• Helps the robot aim and score more consistently\n• Gives drivers clear controls instead of button chaos\n• Uses cameras and sensors so the robot “sees” the field\n• When drivers trust the robot, they can play smarter",
            "notes": "Software = helpful tools for the human drivers.",
        },
        {
            "kicker": "THE BIG IDEA",
            "title": "Why Enzo is the best",
            "body": "1. Built for the whole Crescendo game, not just one trick\n2. The team rebuilt it when they needed to\n3. Focuses on grabbing and controlling Notes\n4. Can score in more than one way\n5. Software helps drivers do the hard parts",
            "notes": "Slow down and hit each of the five points.",
        },
        {
            "kicker": "WRAP-UP",
            "title": "Enzo is best because of how it was built.",
            "body": "Clear plan. Willing to improve. Easy to drive. Ready for the whole game.\n\nQuestions?",
            "notes": "End here. Invite simple questions about design or the game.",
        },
    ]

    for i, data in enumerate(slides, 1):
        s = prs.slides.add_slide(blank)
        fill_bg(s)
        add_text(s, 0.7, 0.45, 12, 0.4, data["kicker"], size=13, bold=True, color=BRASS)
        add_text(s, 0.7, 1.0, 12, 1.4, data["title"], size=32, bold=True, color=CREAM)
        add_text(s, 0.7, 2.6, 11.8, 3.8, data["body"], size=20, color=CREAM)
        add_text(
            s,
            0.7,
            6.9,
            12,
            0.35,
            f"ENZO  ·  FRC Crescendo  ·  Team 1577 Steampunk  ·  {i}/11",
            size=11,
            color=MUTED,
        )
        notes(s, data["notes"])

    out = Path(__file__).resolve().parent / "Enzo_Crescendo_GoogleSlides.pptx"
    prs.save(out)
    return out


if __name__ == "__main__":
    print(build())
